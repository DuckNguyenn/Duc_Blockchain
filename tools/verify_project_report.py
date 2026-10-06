"""Check document structure, saved metrics, provenance and rendered page layout."""
from pathlib import Path
import hashlib
import json
import zipfile
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "report"
SCRATCH = ROOT / "_work_extract" / "report_review_20261005"
docx = OUT / "Bao_cao_blockchain_da_chinh_sua.docx"
pdf = docx.with_suffix(".pdf")
manifest = json.loads((OUT / "review_manifest.json").read_text(encoding="utf-8"))
source = Path(manifest["source_docx_unchanged"])
assert hashlib.sha256(source.read_bytes()).hexdigest() == manifest["sha256"][str(source)]
with zipfile.ZipFile(docx) as z:
    assert z.testzip() is None
    text_xml = z.read("word/document.xml").decode("utf-8")
    assert "TOC" in text_xml
d = Document(docx)
text = "\n".join([p.text for p in d.paragraphs] + [cell.text for table in d.tables for row in table.rows for cell in row.cells])
for obsolete in ["cảm biến sinh hiệu", "huấn luyện liên kết", "chia thưởng", "……………………", "[Bổ sung confusion"]:
    assert obsolete not in text, obsolete
for expected in ["91,14%", "0,8937", "73,86%", "227", "24/24", "6/6", "10/10"]:
    assert expected in text, expected
for name in ["Nguyễn Trần Minh Đức", "Đỗ Thanh Sơn", "Nguyễn Tuấn Kiệt", "Huỳnh Trần Tiến Thịnh", "Nguyễn Hữu Quý", "Huỳnh Thế Thiện"]:
    assert name in text, name
results = json.loads((OUT / "ket_qua_notebook_9.json").read_text(encoding="utf-8"))
assert sum(part["rows"] for part in results["split"].values()) == 1874
assert results["report"]["test_int8"]["accuracy"] == 422 / 463
assert results["report"]["float_int8_agreement"] == 462 / 463

pages = []
with pymupdf.open(pdf) as rendered:
    full = []
    for i, page in enumerate(rendered):
        txt = page.get_text()
        full.append(f"[PAGE {i + 1}]\n{txt}")
        lines = [line.strip() for line in txt.splitlines() if line.strip()]
        assert len(txt.strip()) > 20, f"Blank page {i+1}"
        outside = [w[4] for w in page.get_text("words") if w[0] < -1 or w[1] < -1 or w[2] > page.rect.width+1 or w[3] > page.rect.height+1]
        assert not outside, (i+1, outside)
        pages.append({"page":i+1,"characters":len(txt),"first_lines":lines[:2],"last_lines":lines[-2:]})
        is_body_chapter = any(line.startswith(s) for line in lines[:3] for s in ["CHƯƠNG", "TÀI LIỆU", "PHỤ LỤC"])
        if i == 0 or is_body_chapter or any(s in txt for s in ["Hình 4.2.", "Hình 4.3.", "Bảng 4.5.", "Bảng 4.6.", "Hình 4.4."]):
            page.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(SCRATCH / f"page_{i+1:02d}.png")
    rendered_text = "\n\n".join(full)
    for error in ["Error! Bookmark", "Error! Reference", "No table of contents entries"]:
        assert error not in rendered_text, error
    (SCRATCH / "revised_pdf_text.txt").write_text(rendered_text, encoding="utf-8")
    manifest["render_verification"] = {"pages":len(rendered),"docx_zip_valid":True,"source_unchanged":True,"no_blank_pages":True,"no_text_outside_page":True}
manifest["sha256"][str(docx)] = hashlib.sha256(docx.read_bytes()).hexdigest()
manifest["sha256"][str(pdf)] = hashlib.sha256(pdf.read_bytes()).hexdigest()
(OUT / "review_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
(SCRATCH / "rendered_pages.json").write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"checks":"passed","pages":pages},ensure_ascii=True,indent=2))
