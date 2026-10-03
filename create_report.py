from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(r'C:\Users\P1 Gen 5\Downloads\Blockchain\Bao_cao_SonarChain_HRC_Safety_Log.docx')
BLUE=RGBColor(23,50,77); TEAL=RGBColor(0,122,140); DARK=RGBColor(35,35,35)
NAVY='17324D'; GRAY='F2F5F7'; PALE='EAF3F7'

def font(r,size=12,bold=False,italic=False,color=DARK,name='Times New Roman'):
    r.font.name=name; r._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); r._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name); r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=color

def para(p,align=None,first=True,before=0,after=6,line=1.35):
    if align is not None:p.alignment=align
    f=p.paragraph_format; f.first_line_indent=Cm(.75) if first else None; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line

def P(text='',align=None,size=12,bold=False,italic=False,color=DARK,first=True,style=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph(); para(p,align,first); r=p.add_run(text); font(r,size,bold,italic,color); return p

def H(text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); para(p,first=False,before={1:16,2:11,3:7}[level],after=6); p.paragraph_format.keep_with_next=True; r=p.add_run(text); font(r,{1:16,2:14,3:12.5}[level],True,False,BLUE); return p

def bullet(text,level=0):
    p=doc.add_paragraph(style='List Bullet' if not level else 'List Bullet 2'); p.paragraph_format.left_indent=Cm(.8+level*.5); p.paragraph_format.first_line_indent=Cm(-.35); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.2; r=p.add_run(text); font(r,11.5); return p

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); x=pr.find(qn('w:shd'))
    if x is None:
        x=OxmlElement('w:shd'); pr.append(x)
    x.set(qn('w:val'),'clear'); x.set(qn('w:color'),'auto'); x.set(qn('w:fill'),fill)

def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for k in ('top','start','bottom','end'):
        x=OxmlElement('w:'+k); x.set(qn('w:w'),'100'); x.set(qn('w:type'),'dxa'); m.append(x)
    pr.append(m)

def borders(cell,color='D7E0E5',sz='3'):
    pr=cell._tc.get_or_add_tcPr(); b=pr.find(qn('w:tcBorders'))
    if b is None:
        b=OxmlElement('w:tcBorders'); pr.append(b)
    for k in ('top','left','bottom','right'):
        x=b.find(qn('w:'+k))
        if x is None:
            x=OxmlElement('w:'+k); b.append(x)
        x.set(qn('w:val'),'single'); x.set(qn('w:sz'),sz); x.set(qn('w:color'),color)
    return b

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; margins(c); shade(c,NAVY); borders(c,'FFFFFF','4');
        if widths:c.width=Cm(widths[i])
        p=c.paragraphs[0]; para(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=0,line=1.1); font(p.add_run(h),10.3,True,False,RGBColor(255,255,255))
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            c=cs[i]; c.text=''; margins(c); shade(c,'FFFFFF' if ri%2==0 else GRAY); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Cm(widths[i])
            p=c.paragraphs[0]; para(p,first=False,after=0,line=1.1); font(p.add_run(str(v)),10.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(1); return t

def note(title,text):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,PALE); margins(c); borders(c,'1B8EA3','8'); p=c.paragraphs[0]; para(p,first=False,after=0); font(p.add_run(title+': '),11,True,False,BLUE); font(p.add_run(text),11); doc.add_paragraph().paragraph_format.space_after=Pt(1)

def code(lines):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F7F8F9'); margins(c); borders(c,'9AAAB5','5'); p=c.paragraphs[0]; para(p,first=False,after=0,line=1.0)
    for i,s in enumerate(lines):
        r=p.add_run(s); font(r,9.1,name='Consolas');
        if i<len(lines)-1:r.add_break()
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def field(p,s):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=s; c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'end'); r._r.extend([a,b,c])

def chapter(n,title):
    p=doc.add_paragraph(); para(p,WD_ALIGN_PARAGRAPH.CENTER,False,before=40,after=8); p.paragraph_format.page_break_before=True; font(p.add_run('CHƯƠNG '+str(n)),18,True,False,BLUE)
    p=doc.add_paragraph(); para(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=24); font(p.add_run(title.upper()),16,True,False,TEAL)

def caption(s):
    p=doc.add_paragraph(style='Caption VN'); para(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=8); font(p.add_run(s),10.5,False,True,RGBColor(90,90,90))

doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(3); sec.right_margin=Cm(2); sec.footer_distance=Cm(1)
for st in [doc.styles['Normal'],doc.styles['Heading 1'],doc.styles['Heading 2'],doc.styles['Heading 3']]:
    st.font.name='Times New Roman'; st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman'); st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
doc.styles['Normal'].font.size=Pt(12); doc.styles['Normal'].font.color.rgb=DARK; doc.styles['Normal'].paragraph_format.line_spacing=1.35; doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for nm,sz,col in [('Heading 1',16,BLUE),('Heading 2',14,TEAL),('Heading 3',12.5,BLUE)]:
    s=doc.styles[nm]; s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=col; s.paragraph_format.keep_with_next=True
cap=doc.styles.add_style('Caption VN',WD_STYLE_TYPE.PARAGRAPH); cap.font.name='Times New Roman'; cap.font.size=Pt(10.5); cap.font.italic=True
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; para(f,WD_ALIGN_PARAGRAPH.CENTER,False,after=0); font(f.add_run('Trang '),10,False,False,RGBColor(90,90,90)); field(f,'PAGE')

# Bìa
P('TRƯỜNG ĐẠI HỌC …………………………………',WD_ALIGN_PARAGRAPH.CENTER,14,True,first=False,color=BLUE); P('KHOA / BỘ MÔN …………………………………',WD_ALIGN_PARAGRAPH.CENTER,13,True,first=False,color=BLUE)
for _ in range(3):P('',first=False)
P('BÁO CÁO ĐỒ ÁN / DỰ ÁN',WD_ALIGN_PARAGRAPH.CENTER,15,True,first=False,color=TEAL)
P('HỆ THỐNG GIÁM SÁT AN TOÀN VÙNG CỘNG TÁC NGƯỜI–ROBOT',WD_ALIGN_PARAGRAPH.CENTER,18,True,first=False,color=BLUE)
P('SỬ DỤNG IoT, AI VÀ BLOCKCHAIN',WD_ALIGN_PARAGRAPH.CENTER,17,True,first=False,color=TEAL)
for _ in range(5):P('',first=False)
for a,b in [('Sinh viên thực hiện','........................................................'),('Mã số sinh viên','........................................................'),('Giảng viên hướng dẫn','........................................................'),('Lớp / khóa','........................................................')]: P(a+': '+b,size=12,first=False)
for _ in range(4):P('',first=False)
P('………, tháng …… năm 2026',WD_ALIGN_PARAGRAPH.CENTER,12,False,False,BLUE,False); doc.add_page_break()

# Lời đầu
H('LỜI CẢM ƠN'); P('Em xin trân trọng cảm ơn quý thầy cô Khoa/Bộ môn …………… đã truyền đạt những kiến thức nền tảng về hệ thống nhúng, Internet vạn vật, trí tuệ nhân tạo, cơ sở dữ liệu và công nghệ blockchain. Những kiến thức này là cơ sở để em hình thành ý tưởng, thiết kế và triển khai hệ thống giám sát an toàn trong vùng cộng tác giữa người và robot.')
P('Em xin gửi lời cảm ơn đặc biệt đến giảng viên hướng dẫn …………… vì đã định hướng phạm vi đề tài, góp ý về kiến trúc hệ thống, phương pháp kiểm thử và cách trình bày báo cáo. Các góp ý về ranh giới giữa quyết định an toàn thời gian thực và lớp ghi nhận bằng chứng trên blockchain đã giúp đề tài tránh những diễn giải vượt quá khả năng của một prototype nghiên cứu.')
P('Do giới hạn về thời gian, thiết bị và dữ liệu thử nghiệm, báo cáo vẫn còn những nội dung cần tiếp tục hoàn thiện, đặc biệt là đánh giá thực nghiệm trên phần cứng thật, kiểm định mô hình AI và các yêu cầu của hệ thống an toàn chức năng. Em rất mong nhận được ý kiến đóng góp để hoàn thiện đề tài trong các giai đoạn tiếp theo. Em xin chân thành cảm ơn.')
H('LỜI MỞ ĐẦU'); P('Trong các không gian làm việc có sự phối hợp giữa con người và robot, việc phát hiện sớm tình huống tiếp cận nguy hiểm và lưu lại dấu vết kiểm toán là một nhu cầu quan trọng. Một hệ thống giám sát phù hợp không chỉ cần đo khoảng cách, phân loại trạng thái và cảnh báo tại chỗ, mà còn cần hỗ trợ truy nguyên sự kiện sau khi sự cố xảy ra.')
P('Đề tài xây dựng prototype SonarChain – HRC Safety Log, kết hợp cảm biến siêu âm HC-SR04, ESP32, gateway Python, mô hình phát hiện bất thường bằng AI, smart contract trên mạng EVM cục bộ và dashboard Web3. Nguyên tắc thiết kế là quyết định dừng khẩn cấp phải được xử lý cục bộ theo hướng fail-safe; blockchain chỉ đóng vai trò lớp cam kết bằng chứng và điều phối workflow, không nằm trong vòng điều khiển dừng thời gian thực.')
P('Báo cáo trình bày cơ sở hình thành đề tài, cơ sở lý thuyết, thiết kế từng tầng, kết quả cần đánh giá và giới hạn còn tồn tại. Các mục trong Chương 4 được để trống hoặc đánh dấu rõ khi chưa có dữ liệu, không suy diễn thành kết quả đã đạt được.')
H('DANH MỤC TỪ VIẾT TẮT'); table(['Từ viết tắt','Diễn giải'],[('AI','Artificial Intelligence – Trí tuệ nhân tạo'),('EVM','Ethereum Virtual Machine'),('ESP32','Vi điều khiển có Wi-Fi/Bluetooth'),('HRC','Human–Robot Collaboration'),('IoT','Internet of Things'),('MQTT','Message Queuing Telemetry Transport'),('RPC','Remote Procedure Call'),('SHA-256','Secure Hash Algorithm 256-bit'),('E-Stop','Emergency Stop – Dừng khẩn cấp'),('UI','User Interface')],[4,12.3])
H('DANH MỤC HÌNH VÀ BẢNG'); P('Danh mục đề xuất gồm: Hình 3.1 kiến trúc tổng thể; Hình 3.2 luồng xử lý telemetry; Hình 3.3 ranh giới on-chain/off-chain; Hình 3.4 vòng đời work permit; Bảng 2.1 ngưỡng phân loại; Bảng 3.1 phân công trách nhiệm; Bảng 4.1 ma trận kiểm thử; Bảng 4.2 kết quả AI. Có thể cập nhật sau khi hoàn thành thí nghiệm.')
p=doc.add_paragraph(); para(p,first=False); font(p.add_run('Mục lục'),16,True,False,BLUE); p=doc.add_paragraph(); para(p,first=False); font(p.add_run('Cập nhật mục lục: nhấp chuột phải trong Microsoft Word và chọn “Update Field”.'),10.5,True,False,RGBColor(100,100,100)); p=doc.add_paragraph(); para(p,first=False); field(p,'TOC \\o "1-3" \\h \\z \\u'); doc.add_page_break()

# Ch1
chapter(1,'Giới thiệu và tổng quan'); H('1.1. Bối cảnh và lý do chọn đề tài',2); P('Trong vùng làm việc chung, khoảng cách giữa người và robot có thể thay đổi nhanh do người vận hành di chuyển, robot thực hiện thao tác hoặc xuất hiện vật thể ngoài dự kiến. Tín hiệu cảnh báo cần được phát hiện gần nguồn đo để giảm phụ thuộc vào mạng. Đồng thời, dữ liệu sự kiện nên được lưu theo cách giúp phát hiện việc sửa đổi sau này và hỗ trợ điều tra.')
P('IoT cung cấp khả năng đo và truyền dữ liệu từ edge; AI có thể hỗ trợ nhận biết mẫu bất thường ngoài ngưỡng cố định; blockchain có thể tạo cam kết bất biến đối với bằng chứng off-chain. Mỗi tầng cần một trách nhiệm riêng, không dùng blockchain như bộ điều khiển robot.')
H('1.2. Bài toán đặt ra',2); P('Bài toán là thiết kế prototype có thể tiếp nhận khoảng cách từ cảm biến, phân loại trạng thái, tạo bản ghi bằng chứng, lưu cam kết lên smart contract và hiển thị trạng thái trên dashboard, đồng thời thể hiện được cả luồng an toàn cục bộ và luồng kiểm toán.')
table(['Yêu cầu','Mô tả'],[('R1 – Đo lường','Thu nhận và chuẩn hóa telemetry từ HC-SR04/ESP32.'),('R2 – Phân loại','Áp dụng ngưỡng fail-safe và hỗ trợ AI anomaly detection.'),('R3 – Bằng chứng','Tạo evidence envelope, canonical JSON và SHA-256.'),('R4 – Blockchain','Ghi digest, device hash, timestamp, severity, E-Stop và reporter.'),('R5 – Workflow','Quản lý permit, phê duyệt, zone entry, handoff, complete.'),('R6 – Hiển thị','Dashboard mô phỏng và tương tác MetaMask.'),('R7 – An toàn','Không phụ thuộc RPC/MQTT/blockchain cho dừng cục bộ.')],[3,13.3])
H('1.3. Mục tiêu đề tài',2); [bullet(x) for x in ['Xây dựng pipeline IoT–AI–Blockchain có ranh giới trách nhiệm rõ ràng.','Thiết kế phân loại SAFE/WARNING/EMERGENCY dựa trên ngưỡng, ưu tiên hard threshold trước AI.','Xây dựng HRCSafetyLog và WorkPermitHandoff cho evidence và workflow.','Xây dựng dashboard SonarChain cho mô phỏng, event và permit.']]
H('1.4. Phạm vi và giới hạn',2); P('Phạm vi gồm HC-SR04, ESP32, gateway Python, dữ liệu CSV/JSON, mô hình phát hiện bất thường và mạng Hardhat local. Dashboard hỗ trợ mô phỏng và kết nối MetaMask. Đề tài chưa phải hệ thống safety-certified, chưa chứng minh cơ cấu E-Stop vật lý và không dùng blockchain để điều khiển động cơ hoặc relay. Ngưỡng 60 cm và 30 cm là tham số prototype, không phải khoảng cách bảo vệ được chứng nhận.')
H('1.5. Đóng góp dự kiến',2); P('Đề tài minh họa kiến trúc kết hợp giám sát edge với lớp bằng chứng blockchain mà không làm sai lệch vai trò của blockchain trong hệ thống an toàn; đồng thời cung cấp firmware, gateway, smart contract, test và dashboard để replay, kiểm thử và mở rộng.')
H('1.6. Bố cục báo cáo',2); P('Chương 1 giới thiệu đề tài; Chương 2 trình bày cơ sở lý thuyết; Chương 3 mô tả thiết kế; Chương 4 dành cho kết quả và đánh giá, phần chưa có số liệu để trống; Chương 5 nêu kết luận và hướng phát triển.')

# Ch2
chapter(2,'Cơ sở lý thuyết'); H('2.1. IoT và xử lý tại edge',2); P('IoT là mô hình thiết bị vật lý có khả năng cảm nhận, xử lý và trao đổi dữ liệu. ESP32 là nút edge: đọc cảm biến, tạo telemetry và phát cảnh báo cục bộ. Quyết định sơ bộ ở edge giúp giảm độ trễ và duy trì hành vi an toàn khi MQTT, RPC hoặc dashboard không hoạt động.')
H('2.2. Đo khoảng cách bằng HC-SR04',2); P('HC-SR04 phát xung siêu âm và đo thời gian phản xạ. Khoảng cách ước lượng theo d = v·t/2. Góc bề mặt, vật liệu, nhiễu và vùng chết có thể gây sai lệch; cảm biến trong đề tài là nguồn telemetry cho prototype, không phải cảm biến an toàn đã chứng nhận. Tín hiệu ECHO 5 V phải được hạ mức trước GPIO ESP32.')
H('2.3. Phân loại fail-safe',2); P('Pipeline áp dụng ngưỡng cứng trước khi gọi AI. Nếu khoảng cách nguy hiểm thì EMERGENCY và yêu cầu dừng; vùng cảnh báo là WARNING; chỉ khi không vi phạm ngưỡng mới xét AI.')
table(['Khoảng cách','Mức độ','Hành vi prototype'],[('d > 60 cm','SAFE','Theo dõi, không yêu cầu dừng.'),('30 cm < d ≤ 60 cm','WARNING','Cảnh báo và ghi nhận.'),('d ≤ 30 cm','EMERGENCY','Yêu cầu E-Stop cục bộ; ghi evidence.')],[4,3.5,8.8])
H('2.4. AI phát hiện bất thường',2); P('Các đặc trưng gồm distance_cm và approach_speed_cm_s. Isolation Forest được huấn luyện chủ yếu trên mẫu SAFE để học vùng hoạt động bình thường; mẫu bị cô lập được xem là bất thường. AI chỉ bổ sung cho ngưỡng, không được hạ mức cảnh báo và không phải hệ thống nhận dạng người hoàn chỉnh.')
H('2.5. Hash và bằng chứng toàn vẹn',2); P('SHA-256 tạo digest cố định từ evidence envelope canonical. Dữ liệu thô không đưa trực tiếp lên chain. Hash chứng minh mối liên hệ giữa bản ghi off-chain và cam kết, nhưng không chứng minh cảm biến đo đúng hay phần cứng đã dừng.')
H('2.6. Blockchain và smart contract',2); P('Blockchain cung cấp sổ cái, thứ tự giao dịch, timestamp block và chữ ký tài khoản gửi. HRCSafetyLog là evidence layer; WorkPermitHandoff là workflow layer. Mạng Hardhat local chain ID 31337 phục vụ phát triển và demo.')
H('2.7. Fail-safe và local E-Stop',2); P('Local E-Stop latch không tự mất đi chỉ vì mẫu sau đó an toàn hơn. Việc gỡ dừng cần người có thẩm quyền kiểm tra. Trạng thái E-Stop on-chain là trạng thái logic/audit, không phải bằng chứng cơ cấu vật lý đã dừng.')
H('2.8. MQTT, Serial và gateway',2); P('ESP32 có thể gửi JSON qua Serial hoặc MQTT. Gateway chuẩn hóa, kiểm tra, tính đặc trưng, phân loại, tạo evidence và đưa vào durable outbox. Khi RPC lỗi, bản ghi giữ trạng thái queued/pending để retry.')

# Ch3
chapter(3,'Thiết kế hệ thống'); H('3.1. Kiến trúc tổng thể',2); P('Hệ thống gồm tầng IoT/edge, truyền telemetry, gateway và AI, evidence off-chain, blockchain, backend/workflow và dashboard. Luồng an toàn cục bộ chạy độc lập với luồng ghi blockchain.')
code(['HC-SR04 → ESP32 → Local E-Stop / cảnh báo cục bộ','                    └→ Serial hoặc MQTT → Python Gateway','                                           ├→ Ngưỡng fail-safe','                                           ├→ AI anomaly detection','                                           ├→ Evidence JSON + SHA-256','                                           └→ Outbox → RPC → HRCSafetyLog','                                                        └→ WorkPermitHandoff','Dashboard SonarChain ← ethers.js / MetaMask ← Blockchain local EVM']); caption('Hình 3.1. Kiến trúc tổng thể đề xuất')
table(['Tầng','Thành phần','Trách nhiệm'],[('IoT/Edge','HC-SR04, ESP32, buzzer','Đo, telemetry, cảnh báo cục bộ.'),('Transport','USB Serial, MQTT/Mosquitto','Truyền telemetry.'),('Gateway/AI','Python pipeline, model','Validate, phân loại, evidence.'),('Evidence','Canonical JSON, SHA-256, outbox','Lưu chi tiết và retry.'),('Blockchain','Hai smart contract','Commitment, quyền, E-Stop, permit.'),('Backend','RPC client, adapters','Kết nối off-chain và chain.'),('Dashboard','HTML/CSS/JS, ethers.js','Mô phỏng, hiển thị, ký giao dịch.')],[2.5,5.3,7.5])
H('3.2. Tầng IoT và firmware',2); P('Firmware đọc HC-SR04, phát JSON gồm device_id, sensor_id, timestamp_ms, distance_cm, state, emergency_stop và trạng thái buzzer. Chân TRIG GPIO5, ECHO GPIO18 qua cầu phân áp, buzzer GPIO23 qua transistor và nút silence GPIO27. Mạch là prototype cảnh báo; nút silence không phải E-Stop safety-rated.')
H('3.3. Tầng truyền telemetry',2); P('Serial phù hợp thử nghiệm USB; serial_reader đọc dòng JSON. MQTT dùng Mosquitto, topic mặc định hrc/telemetry/# và QoS 1. Payload thiếu distance_cm, không parse được hoặc có khoảng cách âm bị bỏ qua.')
H('3.4. Gateway và backend',2); P('process_row chuẩn hóa timestamp/distance, tạo đặc trưng, gọi classify, tạo event_id ổn định, xây evidence envelope và enqueue outbox. Backend prototype nằm trong các module Python và RPC adapter; khi triển khai thật cần API service, kho dữ liệu, xác thực, giám sát và reconcile transaction.')
H('3.5. Tầng AI',2); P('Feature engineering tạo distance_cm và approach_speed_cm_s. train_model.py nạp CSV, chia train/test, dùng mẫu SAFE để fit Isolation Forest và ghi model/metrics. AI chỉ bổ sung cho hard threshold.')
table(['Đầu vào','Xử lý','Đầu ra'],[('distance_cm, dt_s','Chuẩn hóa, tính tốc độ','Đặc trưng'),('Đặc trưng','Isolation Forest / model lai','Anomaly, confidence'),('AI + ngưỡng','Ưu tiên hard threshold','SAFE/WARNING/EMERGENCY')],[4.2,6.6,4.5])
H('3.6. Evidence off-chain',2); P('Evidence envelope v1 gồm event_id, device_id, sensor_id, measured_at, received_at, distance_cm, severity, emergency_stop, confidence, policy_version và model_version. JSON được canonicalize và băm; bản ghi lưu trong data/evidence_outbox để retry.')
code(['{','  "schema": "sonarchain.evidence.v1",','  "event_id": "<digest>", "device_id": "HRC-ESP32-01",','  "measured_at": "<UTC>", "distance_cm": 24.6,','  "severity": "EMERGENCY", "emergency_stop": true,','  "policy_version": "thresholds.v60-30"','}'])
H('3.7. HRCSafetyLog',2); P('Contract lưu hash/evidenceHash, deviceIdHash, timestamp, severity, emergencyStop, reporter, evidenceSchema và recordedAt. Contract có allowlist reporter, chống ghi trùng event/evidence, deviceLocked và emergencyStopByDevice. Khi emergencyStop true, thiết bị bị khóa logic; clearEmergencyStop chỉ owner gọi được. Event SAFE sau đó không tự xóa latch.')
H('3.8. WorkPermitHandoff',2); P('Permit có vòng đời NONE, PENDING, APPROVED, ACTIVE, COMPLETED hoặc REVOKED. Requester tạo permit; supervisor phê duyệt; gateway ghi zone entry; participant xác nhận handoff; worker/requester/owner đóng permit. Đây là workflow Web3, không thay thế E-Stop.')
H('3.9. Dashboard SonarChain',2); P('Dashboard hỗ trợ simulation, theo dõi khoảng cách, SAFE/WARNING/EMERGENCY, tạo event hash, kết nối MetaMask và ghi event. Panel permit hỗ trợ tạo, approve, entry, handoff và complete. Sự kiện mô phỏng nằm trong bộ nhớ trình duyệt; nhãn LOCKED không chứng minh motor đã dừng.')
H('3.10. Ranh giới on-chain/off-chain',2); code(['Raw telemetry → Gateway → Evidence JSON → SHA-256 → On-chain commitment','On-chain: digest, device hash, measuredAt, severity, E-Stop, schema, reporter, recordedAt','Off-chain: telemetry chi tiết, confidence, model/policy, evidence file, retry state']); caption('Hình 3.2. Ranh giới dữ liệu on-chain và off-chain')
H('3.11. Thiết kế kiểm thử',2); table(['Nhóm','Trường hợp','Tiêu chí'],[('Gateway','SAFE, WARNING, EMERGENCY, SENSOR_FAULT, dữ liệu sai','Phân loại/log/hash đúng.'),('AI','Train/test và mẫu bất thường','Có model/metrics; không hạ threshold.'),('HRCSafetyLog','Event, trùng hash, quyền, latch','Đúng/revert đúng.'),('WorkPermit','Permit và handoff','Đúng vòng đời/quyền.'),('Tích hợp','Outbox, RPC lỗi, retry','Có trạng thái pending/confirmed.'),('Phần cứng','Khoảng cách, góc, vật liệu','Có sai số/độ trễ.')],[3,8,7.3])

# Ch4
chapter(4,'Kết quả và đánh giá'); note('Hướng dẫn điền chương','Chỉ điền số liệu sau khi chạy đúng kịch bản; ghi cấu hình, phiên bản mã nguồn, thời gian đo và tệp dữ liệu. Không dùng giá trị minh họa làm kết quả chính thức.')
H('4.1. Môi trường thực nghiệm',2); P('[Để trống – bổ sung OS, Python/Node.js, ESP32, cảm biến, Hardhat, chain ID, broker, model và dashboard.]'); table(['Hạng mục','Thông tin cần bổ sung'],[('Phần cứng','ESP32, HC-SR04, buzzer, transistor, cầu phân áp, nguồn.'),('Phần mềm','OS, Python, Node/npm, thư viện, Hardhat, browser, MetaMask.'),('Mạng','RPC URL, chain ID, contract address, broker/topic.'),('Dữ liệu','Số file, số mẫu, số mẫu từng lớp, thời lượng, cách gán nhãn.'),('Phiên bản','Commit/tag repository và model artifact.')],[4,12.3])
H('4.2. Kết quả phần cứng và telemetry',2); P('[Để trống – mô tả nạp firmware, Serial/MQTT, tần suất mẫu, độ trễ và timeout.]'); table(['Chỉ số','Giá trị đo','Ghi chú'],[('Tần suất mẫu','…………','…………'),('Độ trễ cảm biến → gateway','…………','…………'),('Payload hợp lệ','…………','…………'),('Sai số khoảng cách','…………','…………'),('Số SENSOR_FAULT','…………','…………')],[5.5,4,6.8])
H('4.3. Phân loại theo ngưỡng',2); P('[Để trống – bổ sung số mẫu, đúng/sai và confusion matrix theo từng khoảng cách.]'); table(['Kịch bản','Số mẫu','Đúng','Sai','Nhận xét'],[('d > 60 cm','…','…','…','…'),('30 < d ≤ 60 cm','…','…','…','…'),('d ≤ 30 cm','…','…','…','…'),('Timeout/payload lỗi','…','…','…','…')],[4,2.2,2.5,2.5,5.1])
H('4.4. Kết quả mô hình AI',2); P('[Để trống – bổ sung mô hình, train/test, precision, recall, F1, confusion matrix, thời gian suy luận và lý do chọn mô hình.]'); table(['Mô hình','Precision','Recall','F1','Thời gian','Ghi chú'],[('Isolation Forest','…','…','…','…','…'),('Mô hình lai/autoencoder','…','…','…','…','…'),('Mô hình chọn','…','…','…','…','Lý do: …')],[4,2.2,2.2,2.2,3.2,3.1])
H('4.5. Kết quả smart contract',2); P('[Để trống – bổ sung npm test, compile/deploy, transaction, event log và trạng thái deviceLocked/E-Stop.]'); table(['Ca kiểm thử','Mong đợi','Thực tế','Trạng thái'],[('SAFE event','Thành công, không khóa','…………','Đạt/Chưa đạt'),('EMERGENCY','Thành công, khóa','…………','Đạt/Chưa đạt'),('Trùng hash','Revert','…………','Đạt/Chưa đạt'),('Reporter sai','Revert','…………','Đạt/Chưa đạt'),('Clear E-Stop','Chỉ owner','…………','Đạt/Chưa đạt'),('Permit workflow','Đúng vòng đời','…………','Đạt/Chưa đạt')],[4,5.2,5.2,2])
H('4.6. Kết quả dashboard',2); P('[Để trống – chèn ảnh simulation và local blockchain; mô tả kết nối ví, ghi event, E-Stop mô phỏng và permit.]')
H('4.7. Đánh giá đạt mục tiêu',2); P('[Để trống – đối chiếu R1–R7 với log, ảnh hoặc tệp dữ liệu; phân loại Đạt, Đạt một phần hoặc Chưa thực hiện.]'); table(['Yêu cầu','Bằng chứng','Đánh giá','Ghi chú'],[(x,'…………','Đạt/…','…………') for x in ['R1','R2','R3','R4','R5','R6','R7']],[2.2,6,3,5.2])
H('4.8. Hạn chế quan sát được',2); P('[Để trống – ghi nhiễu cảm biến, mất MQTT, RPC timeout, lệch timestamp, UI không đồng bộ hoặc khác biệt mô phỏng/phần cứng.]')

# Ch5
chapter(5,'Kết luận và hướng phát triển'); H('5.1. Kết luận',2); P('Đề tài đã xây dựng prototype SonarChain – HRC Safety Log theo kiến trúc kết hợp IoT, AI và blockchain. Hệ thống lấy telemetry từ HC-SR04/ESP32, xử lý tại gateway, áp dụng ngưỡng fail-safe, hỗ trợ AI anomaly detection, tạo evidence off-chain và ghi digest cùng metadata lên smart contract. Dashboard cung cấp lớp quan sát và tương tác với HRCSafetyLog và WorkPermitHandoff.')
P('Điểm quan trọng là phân tách trách nhiệm: ESP32/gateway xử lý an toàn cục bộ; outbox bảo vệ bằng chứng khi blockchain tạm thời không sẵn sàng; blockchain lưu commitment và workflow; dashboard không phải bộ điều khiển E-Stop. Đây là nền tảng nghiên cứu và trình diễn, chưa phải hệ thống an toàn được chứng nhận. [Bổ sung sau khi có kết quả: nêu mục tiêu đạt, số liệu nổi bật và mục tiêu chưa đạt.]')
H('5.2. Hạn chế',2); [bullet(x) for x in ['HC-SR04 nhạy với góc, vật liệu và môi trường; số cảm biến hạn chế.','Mạng blockchain local và dashboard mô phỏng chưa phản ánh đầy đủ môi trường thật.','Dữ liệu AI chưa đủ để kết luận khả năng tổng quát.','Chưa tích hợp cơ cấu E-Stop vật lý safety-rated và chưa có chứng nhận chức năng.','Còn khoảng trống về schema, verifier, quyền UI, retention và reconcile transaction.']]
H('5.3. Hướng phát triển',2); [bullet(x) for x in ['Bổ sung ToF/lidar/safety scanner và sensor fusion.','Thiết kế E-Stop vật lý độc lập, watchdog, gateway dự phòng, kiểm thử mất mạng/mất nguồn.','Chuẩn hóa evidence schema, ký firmware và xác thực nguồn telemetry.','Bổ sung backend, kho evidence, API verifier và retry/reconcile.','Đánh giá AI bằng dataset lớn hơn, cross-validation, false negative và drift.','Cân nhắc blockchain permissioned, batch/Merkle root và cơ chế bảo mật phù hợp.','Ràng buộc telemetryEventHash với SafetyLog và hoàn thiện policy hạn permit/handoff.']]
H('5.4. Công việc tiếp theo',2); table(['Giai đoạn','Công việc','Đầu ra'],[('1','Đo đạc và điền Chương 4','Dataset, log, ảnh, bảng.'),('2','Chuẩn hóa evidence/verifier','Schema, hash vector, test.'),('3','Kiểm thử lỗi tích hợp','Báo cáo timeout/retry/mất mạng.'),('4','Đánh giá an toàn phần cứng','Thiết kế E-Stop độc lập.'),('5','Đánh giá AI','Metrics, threshold, model card.')],[2,8,7.3])

# Phụ lục
chapter('A','Hướng dẫn chạy và tái lập demo'); H('A.1. Dashboard mô phỏng',2); code(['cd "C:\\Users\\P1 Gen 5\\Downloads\\Blockchain"','py -m http.server 8080 -d web3','Mở http://localhost:8080 và chọn “Chạy mô phỏng”.']); H('A.2. Smart contract local',2); code(['cd contracts','npm.cmd install','npm.cmd run compile','npm.cmd test','npm.cmd run node','npm.cmd run deploy']); H('A.3. Replay CSV',2); code(['python -m iot_code.gateway.pipeline --csv ai_model/data/raw/son_dungyen_30_DANGER.csv --limit 10']); H('A.4. Thông tin cần lưu',2); table(['Thông tin','Giá trị'],[(x,'……………………') for x in ['Thời gian chạy','Commit/tag','Thiết bị/model','Cấu hình ngưỡng','RPC/chain','Tệp log/evidence','Người thực hiện']],[6,11.3])
chapter('B','Thuật ngữ và ranh giới trách nhiệm'); table(['Thuật ngữ','Định nghĩa'],[('Telemetry','Mẫu khoảng cách cùng thông tin nguồn và thời điểm.'),('Safety event','Kết quả đánh giá telemetry theo chính sách an toàn.'),('Off-chain evidence','Bản ghi chi tiết ngoài blockchain.'),('Evidence commitment','Dấu cam kết của bản ghi off-chain.'),('Local E-Stop latch','Trạng thái dừng khẩn cấp tại hệ thống cục bộ.'),('On-chain E-Stop state','Trạng thái logic trên chain, không chứng minh cơ cấu đã dừng.'),('Reporter','Tài khoản được phép ghi safety event.'),('Supervisor','Tài khoản phê duyệt/thu hồi permit.'),('Gateway','Thành phần tiếp nhận telemetry và ghi nhận sự kiện/vào vùng.'),('Work permit','Quyền làm việc có thời hạn cho worker/task/zone.'),('Handoff','Xác nhận nhiệm vụ người–robot đã bàn giao.')],[4,13.3])

doc.save(OUT); print(OUT)
