from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'C:\Users\P1 Gen 5\Downloads\Blockchain\Bao_cao_SonarChain_HRC_Safety_Log.docx')
BLUE = RGBColor(20, 40, 72)
BLACK = RGBColor(25, 25, 25)
GRAY = 'F2F2F2'
NAVY = '18304F'


def set_font(run, size=12, bold=False, italic=False, color=BLACK, name='Times New Roman'):
    run.font.name=name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=color

def fmt(p, align=None, first=True, before=0, after=6, line=1.3):
    if align is not None: p.alignment=align
    pf=p.paragraph_format
    pf.first_line_indent=Cm(.7) if first else None
    pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=line

def add_p(text='', align=None, size=12, bold=False, italic=False, color=BLACK, first=True, style=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph()
    fmt(p,align,first)
    set_font(p.add_run(text),size,bold,italic,color)
    return p

def heading(text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}')
    fmt(p, first=False, before={1:14,2:9,3:5}[level], after=5)
    p.paragraph_format.keep_with_next=True
    set_font(p.add_run(text), {1:15,2:13,3:12}[level], True, False, BLUE)
    return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Cm(.9); p.paragraph_format.first_line_indent=Cm(-.35); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.2; set_font(p.add_run(text),11.5); return p

def shade(cell, fill):
    pr=cell._tc.get_or_add_tcPr(); sh=pr.find(qn('w:shd'))
    if sh is None: sh=OxmlElement('w:shd'); pr.append(sh)
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY)
        if widths: c.width=Cm(widths[i])
        p=c.paragraphs[0]; fmt(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=0,line=1.05)
        for r in p.runs: set_font(r,10,True,False,RGBColor(255,255,255))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); shade(cells[i], 'FFFFFF' if ri%2==0 else GRAY)
            if widths: cells[i].width=Cm(widths[i])
            p=cells[i].paragraphs[0]; fmt(p,first=False,after=0,line=1.08)
            for r in p.runs: set_font(r,9.8)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def code(lines):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F4F5F6'); c.text=''
    p=c.paragraphs[0]; fmt(p,first=False,after=0,line=1.0)
    for i,line in enumerate(lines):
        r=p.add_run(line); set_font(r,9,False,False,BLACK,'Consolas')
        if i<len(lines)-1: r.add_break()
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def caption(text):
    p=doc.add_paragraph(style='Caption VN'); fmt(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=7); set_font(p.add_run(text),10,False,True,RGBColor(80,80,80))

def page_field(p, instruction):
    r=p.add_run(); b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin'); i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=instruction; e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end'); r._r.extend([b,i,e])

def chapter(num,title):
    p=doc.add_paragraph(); fmt(p,WD_ALIGN_PARAGRAPH.CENTER,False,before=28,after=5); p.paragraph_format.page_break_before=True; set_font(p.add_run('CHƯƠNG '+str(num)),17,True,False,BLUE)
    p=doc.add_paragraph(); fmt(p,WD_ALIGN_PARAGRAPH.CENTER,False,after=18); set_font(p.add_run(title.upper()),15,True,False,BLUE)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.8); sec.right_margin=Cm(2.0); sec.header_distance=Cm(0.8); sec.footer_distance=Cm(.8)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman'); normal._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman'); normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.3; normal.paragraph_format.space_after=Pt(5)
for nm,sz in [('Heading 1',15),('Heading 2',13),('Heading 3',12)]:
    s=doc.styles[nm]; s.font.name='Times New Roman'; s._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman'); s._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman'); s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=BLUE
cap=doc.styles.add_style('Caption VN',WD_STYLE_TYPE.PARAGRAPH); cap.font.name='Times New Roman'; cap.font.size=Pt(10); cap.font.italic=True
footer=sec.footer.paragraphs[0]; fmt(footer,WD_ALIGN_PARAGRAPH.CENTER,False,after=0); set_font(footer.add_run('Trang '),10,False,False,RGBColor(90,90,90)); page_field(footer,'PAGE')

# Cover: similar to sample
add_p('TRƯỜNG ĐẠI HỌC …………………………………',WD_ALIGN_PARAGRAPH.CENTER,14,True,color=BLUE,first=False)
add_p('KHOA ĐIỆN – ĐIỆN TỬ',WD_ALIGN_PARAGRAPH.CENTER,13,True,color=BLUE,first=False)
add_p('BỘ MÔN KỸ THUẬT MÁY TÍNH',WD_ALIGN_PARAGRAPH.CENTER,13,True,color=BLUE,first=False)
for _ in range(4): add_p('',first=False)
add_p('MÔN HỌC: BLOCKCHAIN VÀ ỨNG DỤNG',WD_ALIGN_PARAGRAPH.CENTER,14,True,color=BLUE,first=False)
add_p('ĐỀ TÀI:',WD_ALIGN_PARAGRAPH.CENTER,14,True,color=BLUE,first=False)
add_p('HỘP ĐEN AN TOÀN VÀ ĐỊNH DANH CHO ROBOT HỢP TÁC',WD_ALIGN_PARAGRAPH.CENTER,16,True,color=BLUE,first=False)
add_p('(SONARCHAIN – HRC SAFETY LOG)',WD_ALIGN_PARAGRAPH.CENTER,14,True,color=BLUE,first=False)
for _ in range(3): add_p('',first=False)
for label,value in [('GVHD','…………………………………'),('Mã môn học','…………………………………'),('Ngành','Công Nghệ Kỹ Thuật Máy Tính'),('Khóa','2023'),('Nhóm thực hiện','Nhóm ………')]: add_p(f'{label:<22}: {value}',size=12,first=False)
add_p('',first=False); add_p('DANH SÁCH THÀNH VIÊN',WD_ALIGN_PARAGRAPH.CENTER,12,True,color=BLUE,first=False)
table(['Họ và tên','MSSV','Hoàn thành'],[('………………………………','…………','……%'),('………………………………','…………','……%'),('………………………………','…………','……%'),('………………………………','…………','……%')],[8,3,4.3])
for _ in range(2): add_p('',first=False)
add_p('Tp. Hồ Chí Minh, tháng …… năm 2026',WD_ALIGN_PARAGRAPH.CENTER,12,False,color=BLUE,first=False); doc.add_page_break()

# Front matter
heading('LỜI CẢM ƠN'); add_p('Trong quá trình học tập và thực hiện báo cáo cuối kỳ của môn học Blockchain và Ứng dụng, nhóm chúng em đã nhận được sự hướng dẫn và hỗ trợ từ quý Thầy Cô cùng các bạn. Những kiến thức về blockchain, smart contract, hệ thống nhúng và trí tuệ nhân tạo là nền tảng để nhóm xây dựng đề tài SonarChain – HRC Safety Log.')
add_p('Nhóm chúng em xin bày tỏ lòng biết ơn đến giảng viên hướng dẫn …………… đã định hướng bài toán, góp ý về kiến trúc IoT – AI – Blockchain, phương pháp kiểm thử và cách trình bày báo cáo. Các góp ý về việc tách luồng dừng an toàn tại edge khỏi luồng ghi bằng chứng trên blockchain giúp nhóm hiểu rõ hơn giới hạn của một prototype nghiên cứu.')
add_p('Nhóm cũng xin cảm ơn các bạn đã trao đổi và hỗ trợ trong quá trình tìm hiểu ESP32, cảm biến HC-SR04, MQTT, Hardhat, MetaMask và các công cụ phát triển liên quan. Do thời gian và kinh nghiệm còn hạn chế, báo cáo vẫn còn những nội dung cần bổ sung, đặc biệt là số liệu đo trên phần cứng và đánh giá định lượng mô hình AI. Nhóm rất mong nhận được góp ý để hoàn thiện đề tài.')
add_p('Nhóm xin chân thành cảm ơn!\n\nNhóm thực hiện đề tài\n(Ký và ghi rõ họ tên)',WD_ALIGN_PARAGRAPH.RIGHT,12,False,False,BLACK,False)
heading('TÓM TẮT'); add_p('Đề tài SonarChain – HRC Safety Log nghiên cứu việc kết hợp IoT, AI và Blockchain để xây dựng một lớp giám sát và truy vết sự kiện an toàn trong không gian cộng tác giữa người và robot. Hệ thống sử dụng cảm biến siêu âm HC-SR04 kết nối với ESP32 để đo khoảng cách, gateway Python để xử lý telemetry, mô hình phát hiện bất thường để hỗ trợ phân tích và smart contract để lưu cam kết bằng chứng.')
add_p('Kiến trúc được chia thành hai luồng. Luồng an toàn tại edge áp dụng ngưỡng fail-safe và có thể kích hoạt yêu cầu E-Stop cục bộ mà không phụ thuộc MQTT, RPC hoặc blockchain. Luồng truy vết tạo evidence envelope off-chain, tính SHA-256, lưu bản ghi vào outbox và ghi digest cùng metadata lên HRCSafetyLog. Smart contract WorkPermitHandoff bổ sung workflow cấp phép làm việc, zone entry và human–robot handoff. Dashboard SonarChain hỗ trợ mô phỏng telemetry, hiển thị sự kiện và ký giao dịch bằng MetaMask khi chạy mạng EVM local.')
add_p('[Kết quả định lượng về IoT, AI, smart contract và end-to-end sẽ được bổ sung sau khi nhóm hoàn thành các kịch bản thực nghiệm. Báo cáo không sử dụng số liệu chưa được đo kiểm làm kết quả chính thức.]',size=12,italic=True,color=RGBColor(80,80,80))

# TOC
p=doc.add_paragraph(); fmt(p,first=False); set_font(p.add_run('Mục lục'),15,True,False,BLUE)
p=doc.add_paragraph(); fmt(p,first=False); page_field(p,'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()
heading('DANH SÁCH HÌNH VẼ'); add_p('Hình 3.1. Sơ đồ kiến trúc tổng thể hệ thống.\nHình 3.2. Luồng dữ liệu telemetry và bằng chứng.\nHình 3.3. Vòng đời work permit và human–robot handoff.\nHình 4.1. [Bổ sung] Kết quả mô phỏng/dashboard.\nHình 4.2. [Bổ sung] Kết quả đánh giá mô hình AI.',first=False)
heading('DANH SÁCH BẢNG'); add_p('Bảng 2.1. Ngưỡng phân loại khoảng cách.\nBảng 3.1. Linh kiện và chân kết nối.\nBảng 3.2. Cấu trúc evidence envelope.\nBảng 3.3. Vai trò trong các smart contract.\nBảng 3.4. Trình tự xử lý một sự cố.\nBảng 4.1. Môi trường thực nghiệm.\nBảng 4.2. Kết quả tầng IoT.\nBảng 4.3. Kết quả AI.\nBảng 4.4. Kết quả smart contract và end-to-end.',first=False)

# Chapter 1
chapter(1,'Giới thiệu tổng quan'); heading('1.1. Giới thiệu bài toán',2); add_p('Trong các nhà máy hiện đại, robot hợp tác có thể làm việc trong cùng không gian với con người để tăng tính linh hoạt của dây chuyền. Tuy nhiên, việc dùng chung vùng làm việc cũng làm tăng yêu cầu về giám sát khoảng cách, cảnh báo và truy vết sau sự cố. Một hệ thống chỉ cảnh báo tại chỗ có thể không cung cấp đủ dữ liệu để phân tích nguyên nhân; ngược lại, một hệ thống chỉ dựa vào mạng hoặc blockchain có thể có độ trễ không phù hợp với quyết định dừng thời gian thực.')
add_p('Đề tài lựa chọn bài toán xây dựng hộp đen an toàn và định danh cho robot hợp tác. Dữ liệu telemetry được thu nhận tại edge, phân loại bằng quy tắc fail-safe và AI, sau đó được đóng gói thành bằng chứng off-chain. Blockchain không lưu dữ liệu thô mà lưu digest và metadata cần thiết để kiểm tra tính toàn vẹn, nguồn ghi và trạng thái logic của thiết bị.')
add_p('Xuất phát từ hai yêu cầu trên, nhóm xây dựng SonarChain – HRC Safety Log, một prototype tích hợp IoT – AI – Blockchain, có thể chạy ở chế độ mô phỏng hoặc kết nối với mạng Hardhat local và MetaMask.')
heading('1.2. Mục tiêu đề tài',2)
for x in ['Xây dựng hệ thống đo khoảng cách dùng ESP32 và cảm biến HC-SR04, có ngưỡng SAFE/WARNING/EMERGENCY.','Thiết kế gateway tiếp nhận Serial/MQTT, kiểm tra dữ liệu, tạo đặc trưng và hỗ trợ phát hiện bất thường bằng AI.','Tạo evidence envelope versioned, canonical JSON và SHA-256 để bảo vệ tính toàn vẹn của bản ghi off-chain.','Xây dựng HRCSafetyLog để lưu commitment, phân quyền reporter và duy trì trạng thái khóa/E-Stop logic.','Xây dựng WorkPermitHandoff để quản lý permit, phê duyệt, zone entry và handoff.','Xây dựng dashboard SonarChain tích hợp MetaMask và chế độ mô phỏng.']: bullet(x)
heading('1.3. Phạm vi và giới hạn',2); add_p('Phạm vi gồm ESP32, HC-SR04, buzzer, gateway Python, dữ liệu CSV/JSON, mô hình Isolation Forest hoặc mô hình lai, hai smart contract Solidity và dashboard HTML/JavaScript. Mạng blockchain dùng Hardhat local cho phát triển. Các ngưỡng 60 cm và 30 cm là tham số demo; hệ thống chưa được chứng nhận an toàn chức năng, chưa chứng minh cơ cấu E-Stop vật lý và không dùng blockchain để điều khiển motor hoặc relay.')

# Ch2
chapter(2,'Cơ sở lý thuyết'); heading('2.1. Tổng quan về CPS và robot hợp tác',2); add_p('Cyber Physical System (CPS) là hệ thống kết hợp giữa quá trình vật lý, cảm biến, bộ điều khiển, mạng truyền thông và phần mềm. Trong đề tài, HC-SR04 và ESP32 thuộc phần vật lý; gateway, AI, evidence service và blockchain thuộc phần tính toán – truyền thông. Robot hợp tác (HRC) là hình thức con người và robot chia sẻ một vùng làm việc, do đó cần theo dõi trạng thái môi trường và xử lý rủi ro theo thời gian.')
heading('2.2. Giám sát khoảng cách và nguyên tắc fail-safe',2); add_p('Nguyên tắc giám sát khoảng cách là khi khoảng cách đo được giảm xuống vùng nguy hiểm, hệ thống phải yêu cầu dừng hoặc chuyển sang trạng thái bảo vệ. Trong prototype, hệ thống dùng ngưỡng cố định để dễ kiểm tra và giải thích. Khi lỗi cảm biến hoặc mất dữ liệu, policy thực tế cần tránh coi mẫu lỗi là SAFE; đây là cơ sở để bổ sung trạng thái SENSOR_FAULT trong các phiên bản sau.')
table(['Khoảng cách','Trạng thái','Hành vi'],[('d > 60 cm','SAFE','Theo dõi, không khóa.'),('30 < d ≤ 60 cm','WARNING','Cảnh báo và ghi nhận tùy chính sách.'),('d ≤ 30 cm','EMERGENCY','Yêu cầu E-Stop cục bộ và có thể khóa logic.')],[4,4,8.3])
heading('2.3. Blockchain, Smart Contract và hàm băm',2); add_p('Blockchain là sổ cái trong đó các giao dịch được xác nhận và sắp xếp theo thứ tự. Smart contract là chương trình Solidity thực thi trên EVM, có thể kiểm tra quyền, điều kiện chuyển trạng thái và phát event. Do dữ liệu trên chain công khai và chi phí ghi cao, hệ thống chỉ lưu digest của evidence, deviceIdHash, thời điểm đo, severity, cờ E-Stop, schema và reporter.')
add_p('SHA-256 được dùng để tạo evidence_hash từ evidence envelope canonical. Trong dashboard và adapter EVM, digest có thể được biểu diễn thành bytes32 bằng quy tắc băm phù hợp. Hash chứng minh bản ghi hiện tại khớp với cam kết đã ghi; hash không chứng minh cảm biến đo đúng hoặc động cơ đã dừng.')
heading('2.4. Isolation Forest',2); add_p('Isolation Forest là thuật toán phát hiện bất thường không giám sát. Thuật toán tạo nhiều cây chia ngẫu nhiên không gian đặc trưng; mẫu khác biệt thường bị cô lập với độ dài đường đi ngắn hơn. Đề tài sử dụng distance_cm và approach_speed_cm_s, huấn luyện chủ yếu trên mẫu SAFE. AI chỉ bổ sung cho hard threshold và không được phép hạ cảnh báo EMERGENCY.')
heading('2.5. MQTT, Serial và evidence envelope',2); add_p('MQTT dùng mô hình publish/subscribe với broker trung gian, phù hợp với thiết bị tài nguyên hạn chế. Serial qua USB là kênh đơn giản để thử nghiệm ESP32. Gateway chuẩn hóa cả hai kênh về cùng một dạng row trước khi xử lý. Evidence envelope là JSON có schema version, giúp các producer và verifier biết cách canonicalize và băm dữ liệu.')

# Ch3
chapter(3,'Thiết kế hệ thống'); heading('3.1. Kiến trúc tổng thể',2); add_p('Hệ thống được tổ chức thành hai luồng song song. Luồng an toàn (fast lane) chạy tại ESP32/gateway, áp dụng ngưỡng cứng và duy trì local E-Stop latch. Luồng truy vết (trust lane) đưa telemetry qua gateway, tạo evidence, lưu outbox và submit digest lên blockchain. Do blockchain có độ trễ và có thể mất kết nối, blockchain không nằm trong điều kiện duy nhất để dừng robot.')
code(['HC-SR04 → ESP32 → Local E-Stop / cảnh báo cục bộ','                    └→ Serial hoặc MQTT → Gateway Python','                                           ├→ Validate + hard threshold','                                           ├→ AI anomaly detection','                                           ├→ Evidence JSON + SHA-256','                                           └→ Outbox → RPC → HRCSafetyLog','                                                        └→ WorkPermitHandoff','Dashboard SonarChain ← ethers.js / MetaMask ← EVM local'])
caption('Hình 3.1: Sơ đồ kiến trúc tổng thể hệ thống')
heading('3.2. Tầng IoT',2); heading('3.2.1. Vai trò',3); add_p('Tầng IoT thu thập khoảng cách, phát trạng thái telemetry và cảnh báo tại edge. ESP32 phải tiếp tục áp dụng logic cục bộ ngay cả khi broker, gateway hoặc blockchain không hoạt động.')
heading('3.2.2. Phần cứng',3); table(['Thiết bị','Kết nối / cấu hình','Vai trò'],[('ESP32 DevKit','Bộ điều khiển edge','Đọc cảm biến và phát JSON.'),('HC-SR04','TRIG GPIO5; ECHO GPIO18 qua cầu phân áp','Đo khoảng cách.'),('Buzzer','GPIO23 qua transistor NPN','Cảnh báo âm thanh.'),('Nút nhấn','GPIO27, INPUT_PULLUP','Silence/acknowledge, không phải E-Stop safety-rated.'),('Nguồn và GND','VIN/5V, mass chung','Cấp nguồn và tham chiếu.')],[4,6.5,5.8])
add_p('ECHO của HC-SR04 thường ở mức 5 V nên phải dùng cầu phân áp trước GPIO ESP32. Firmware hiện có thể phát các trường device_id, sensor_id, timestamp_ms, distance_cm, state, emergency_stop, buzzer_on và seq.')
heading('3.2.3. Thuật toán phân loại',3); code(['if distance_cm <= 30:       severity = EMERGENCY; emergency_stop = true','elif distance_cm <= 60:     severity = WARNING; emergency_stop = false','else:                         severity = SAFE; emergency_stop = false','Sau hard threshold, nếu có model hợp lệ mới xét anomaly/confidence để bổ sung WARNING.'])
heading('3.2.4. Gói tin telemetry',3); code(['{"device_id":"ESP32-HRC-01","sensor_id":"HC-SR04",',' "timestamp_ms":1234,"distance_cm":24.7,',' "state":"EMERGENCY","emergency_stop":true,"seq":12}'])
heading('3.3. Tầng AI',2); heading('3.3.1. Vai trò',3); add_p('AI hỗ trợ phát hiện mẫu bất thường như tín hiệu nhảy đột ngột hoặc tốc độ tiếp cận không giống vùng hoạt động bình thường. AI không nằm trong vòng dừng cục bộ; hard threshold có độ ưu tiên cao hơn.')
heading('3.3.2. Tập dữ liệu',3); add_p('Repository có các file CSV raw gồm mẫu SAFE, DANGER/EMERGENCY, APPROACHING, RETREATING và các kịch bản vật thể. Mỗi bản ghi có timestamp, elapsed_s, distance_cm và label. Khi hoàn thiện báo cáo, cần bổ sung số lượng thực tế và cách chia train/test theo phiên đo để tránh rò rỉ dữ liệu.')
table(['Tập dữ liệu','Số bản ghi','SAFE','Bất thường','Ghi chú'],[('Train','……','……','……','Bổ sung sau khi chạy.'),('Validation','……','……','……','Bổ sung sau khi chạy.'),('Test','……','……','……','Bổ sung sau khi chạy.')],[3.5,3,3,3,4.8])
heading('3.3.3. Kiến trúc mô hình',3); add_p('Pipeline cơ bản dùng Isolation Forest với StandardScaler và huấn luyện trên mẫu SAFE. Một số artifact nâng cao hỗ trợ classifier, Isolation Forest và autoencoder với cơ chế fusion. Kết quả model được lưu bằng joblib cùng thresholds và metrics.')
heading('3.3.4. Triển khai',3); add_p('Gateway nạp model nếu file tồn tại; nếu không có model, pipeline vẫn thực hiện phân loại theo ngưỡng. Confidence, model_version và policy_version được đưa vào evidence để có thể truy nguyên phiên bản quyết định.')
heading('3.4. Tầng Blockchain',2); heading('3.4.1. Vai trò',3); add_p('HRCSafetyLog là lớp evidence commitment: lưu bằng chứng tối thiểu để đối chiếu bản ghi off-chain. WorkPermitHandoff là lớp workflow: lưu quyền làm việc, phê duyệt, zone entry và handoff. Cả hai contract không đọc cảm biến và không điều khiển động cơ.')
heading('3.4.2. Phân quyền',3); table(['Vai trò','Được phép','Ý nghĩa'],[('Owner','Quản lý reporter; clear E-Stop; quản trị contract','Tài khoản triển khai/điều hành.'),('Reporter','recordEvent/recordEvidence','Đưa safety event lên log.'),('Supervisor','Approve/revoke permit','Phê duyệt workflow.'),('Gateway','Record zone entry','Ghi nhận worker vào zone dựa trên telemetry.'),('Worker/Requester','Create permit; confirm handoff; complete theo điều kiện','Tham gia workflow.')],[3.2,6.5,6.6])
heading('3.4.3. Máy trạng thái và hàm',3); add_p('HRCSafetyLog duy trì eventExists, evidenceExists, deviceLocked và emergencyStopByDevice. Khi ghi EMERGENCY, device bị khóa logic; clearEmergencyStop chỉ owner gọi được. WorkPermitHandoff chuyển PENDING → APPROVED → ACTIVE → COMPLETED, hoặc REVOKED theo điều kiện.')
table(['Contract / hàm','Quyền','Chức năng'],[('recordEvidence','Reporter','Ghi digest, device hash, timestamp, severity, E-Stop, schema.'),('clearEmergencyStop','Owner','Gỡ latch logic trên chain.'),('createPermit','Bất kỳ requester','Tạo permit có thời hạn.'),('approvePermit','Supervisor','Phê duyệt permit.'),('recordZoneEntry','Gateway','Đưa permit sang ACTIVE.'),('confirmHandoff','Worker/requester/supervisor','Xác nhận bàn giao.'),('completePermit','Participant hợp lệ','Đóng permit.')],[4.8,4.2,7.3])
heading('3.4.4. Evidence và chống ghi trùng',3); add_p('Contract từ chối eventHash/evidenceHash đã tồn tại, schema không hỗ trợ, timestamp rỗng hoặc cờ emergency không phù hợp severity. Cơ chế này chống ghi trùng cùng một commitment; việc xác minh nội dung off-chain vẫn cần verifier băm lại evidence và đối chiếu transaction.')
heading('3.5. Backend và Dashboard',2); heading('3.5.1. Vai trò',3); add_p('Backend prototype gồm gateway pipeline, MQTT subscriber, serial reader, evidence outbox và blockchain client. Dashboard HTML/CSS/JavaScript dùng ethers.js và MetaMask để mô phỏng, tạo event local, ghi event, quản trị role, E-Stop logic và work permit.')
heading('3.5.2. Sơ đồ tuần tự xử lý sự cố',3); table(['Bước','Thành phần','Việc xảy ra'],[('1','HC-SR04 → ESP32','Đo distance.'),('2','ESP32/gateway','Áp hard threshold; local E-Stop nếu nguy hiểm.'),('3','ESP32 → Serial/MQTT','Gửi telemetry.'),('4','Gateway','Validate, tạo đặc trưng, chạy AI nếu có.'),('5','Gateway → Outbox','Tạo evidence envelope và SHA-256.'),('6','Outbox → HRCSafetyLog','Submit digest/metadata; lỗi RPC giữ pending.'),('7','Dashboard','Hiển thị severity, hash, tx và trạng thái.'),('8','Owner/Supervisor','Clear E-Stop hoặc hoàn tất workflow sau kiểm tra.')],[2,5,10.3])

# Ch4
chapter(4,'Kết quả'); add_p('[Các mục trong chương này được giữ theo bố cục báo cáo mẫu. Nội dung có dấu “……” cần được thay bằng số liệu thực tế sau khi chạy thí nghiệm; không coi đây là kết quả đã xác nhận.]',size=12,italic=True,color=RGBColor(80,80,80))
heading('4.1. Môi trường và kịch bản',2); table(['Thành phần','Cấu hình thực nghiệm'],[('Máy tính','Windows 11; Python ………; Node.js ………'),('Blockchain','Hardhat; chain ID 31337; Solidity ………'),('MQTT','Mosquitto ………; port 1883 hoặc ………'),('Firmware','ESP32 DevKit; Arduino IDE/PlatformIO ………'),('AI','scikit-learn ………; model artifact ………'),('Dashboard','Trình duyệt ………; ethers.js ………; MetaMask ………')],[4.5,12])
add_p('Kịch bản đề xuất gồm: khoảng cách SAFE ổn định; vật thể tiến vào WARNING; khoảng cách ≤ 30 cm tạo EMERGENCY; payload không hợp lệ hoặc timeout; RPC lỗi rồi retry; workflow permit đầy đủ từ create đến complete.')
heading('4.2. Kết quả IoT',2); table(['Hạng mục','Kết quả'],[('Đọc HC-SR04','………………………………'),('Chu kỳ lấy mẫu','………………………………'),('Tỷ lệ payload hợp lệ','………………………………'),('Độ trễ sensor → gateway','………………………………'),('Tình huống timeout','………………………………'),('Sai số đo khoảng cách','………………………………')],[6,10.3])
heading('4.3. Kết quả AI',2); table(['Mô hình','Precision','Recall','F1-score','Ghi chú'],[('Isolation Forest','……','……','……','……'),('Mô hình lai/autoencoder','……','……','……','……'),('Mô hình được chọn','……','……','……','Lý do: ………')],[4,2.8,2.8,2.8,4.9]); add_p('[Bổ sung confusion matrix, phân bố anomaly score và phân tích false negative/false positive nếu có.]')
heading('4.4. Kết quả Smart Contract',2); table(['Ca kiểm thử','Kết quả mong đợi','Kết quả thực tế'],[('Record SAFE','Thành công; không khóa device','……………………'),('Record EMERGENCY','Thành công; khóa device','……………………'),('Duplicate hash','Revert event already recorded','……………………'),('Reporter trái phép','Revert not reporter','……………………'),('Clear E-Stop','Chỉ owner; latch false','……………………'),('Permit workflow','Đúng trạng thái và quyền','……………………')],[4,7,6.3])
heading('4.5. Kết quả end-to-end',2); table(['Kịch bản','Telemetry','AI/gateway','Blockchain','Dashboard'],[('SAFE','……','……','……','……'),('WARNING','……','……','……','……'),('EMERGENCY','……','……','……','……'),('RPC mất kết nối','……','Outbox pending','……','……'),('Permit + handoff','……','……','……','……')],[3.5,3,3.4,3.4,3.0]); add_p('[Chèn ảnh chụp dashboard, transaction hash, event log và tệp evidence tương ứng. Các ảnh/số liệu cần ghi kèm thời gian và commit mã nguồn.]')
heading('4.6. Đánh giá mục tiêu',2); table(['Mục tiêu','Bằng chứng','Đánh giá'],[(f'R{i}','……………………','Đạt / Đạt một phần / Chưa đạt') for i in range(1,8)],[2.2,8,7])

# Ch5
chapter(5,'Kết luận và hướng phát triển'); heading('5.1. Kết luận',2); add_p('Đề tài đã thiết kế một prototype HRC Safety Log theo hướng tích hợp IoT, AI và Blockchain. Hệ thống sử dụng ESP32 và HC-SR04 để tạo telemetry, gateway để áp dụng chính sách fail-safe và hỗ trợ AI, evidence outbox để lưu dữ liệu chi tiết, HRCSafetyLog để ghi commitment và WorkPermitHandoff để quản lý workflow người–robot.')
add_p('Kiến trúc hai luồng là nguyên tắc cốt lõi: quyết định dừng phải nằm tại edge/gateway, còn blockchain cung cấp lớp audit và coordination. Cách phân tách này giúp hệ thống không bị phụ thuộc vào độ trễ hoặc tình trạng sẵn sàng của mạng blockchain. [Bổ sung sau thực nghiệm: các mục tiêu đạt được và số liệu chính.]')
heading('5.2. Hướng phát triển',2)
for x in ['Bổ sung cảm biến ToF, lidar hoặc safety scanner và sensor fusion.','Thiết kế E-Stop vật lý độc lập, watchdog, gateway dự phòng và kiểm thử mất mạng/mất nguồn.','Chuẩn hóa evidence schema xuyên suốt firmware, gateway, contract và dashboard; thêm verifier đối chiếu chain.','Bổ sung backend service, database, xác thực, retention và retry/reconcile transaction.','Đánh giá AI bằng dataset lớn hơn, chia theo phiên đo, phân tích false negative và drift.','Chuyển sang blockchain permissioned hoặc batch/Merkle root khi cần giảm số transaction.','Ràng buộc telemetryEventHash của zone entry với SafetyLog và hoàn thiện policy hết hạn permit.']: bullet(x)
heading('5.3. Hạn chế',2); add_p('Prototype chưa phải hệ thống safety-certified; HC-SR04 có giới hạn về góc đo, vật liệu và nhiễu; dữ liệu AI chưa đủ để khẳng định khả năng tổng quát; dashboard mô phỏng giữ state trong bộ nhớ; trạng thái LOCKED trên chain không chứng minh cơ cấu vật lý đã dừng.')

# Appendix
chapter('A','Hướng dẫn chạy và tái lập demo'); heading('A.1. Dashboard mô phỏng',2); code(['cd "C:\\Users\\P1 Gen 5\\Downloads\\Blockchain"','py -m http.server 8080 -d web3','Mở http://localhost:8080 → Chạy mô phỏng'])
heading('A.2. Smart contract local',2); code(['cd contracts','npm.cmd install','npm.cmd run compile','npm.cmd test','npm.cmd run node','npm.cmd run deploy'])
heading('A.3. Replay dữ liệu',2); code(['python -m iot_code.gateway.pipeline --csv ai_model/data/raw/son_dungyen_30_DANGER.csv --limit 10'])
heading('A.4. Dữ liệu cần lưu khi chạy',2); table(['Thông tin','Giá trị'],[(x,'………………………………') for x in ['Thời gian','Commit/tag','Model/config','RPC/contract','File evidence','Người thực hiện']],[5,12.3])

# settings update fields
try:
    settings=doc.settings._element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
except Exception: pass
doc.save(OUT); print(OUT)
